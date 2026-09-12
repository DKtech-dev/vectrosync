"""
Enterprise Technical Presentation Builder (build_full_presentation.py)
Asset: Well #14, Baghewala Heavy Oil Asset, Bikaner-Nagaur Basin, Rajasthan | Operator: Oil India Limited
Model: OIL-BAGHEWALA-EOR-V2
"""

import sys
import os
import pptx
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

# Color Palette Definitions
C_NAVY = RGBColor(15, 23, 42)       # #0F172A - Primary Navy
C_SLATE = RGBColor(71, 85, 105)     # #475569 - Muted Slate
C_LIGHT_SLATE = RGBColor(100, 116, 139) # #64748B
C_BLUE = RGBColor(2, 132, 199)      # #0284C7 - Industrial Cyan/Blue
C_EMERALD = RGBColor(5, 150, 105)   # #059669 - Safe State Emerald
C_CRIMSON = RGBColor(220, 38, 38)   # #DC2626 - Failure Alert Crimson
C_CARD_BG = RGBColor(248, 250, 252) # #F8FAFC - Card Surface
C_BORDER = RGBColor(226, 232, 240)  # #E2E8F0 - 1px Border
C_WHITE = RGBColor(255, 255, 255)
C_DARK_BG = RGBColor(14, 17, 23)    # #0E1117 - Live UI Surface
C_AMBER = RGBColor(217, 119, 6)     # #D97706
C_LIGHT_GREEN = RGBColor(209, 250, 229)
C_LIGHT_RED = RGBColor(254, 226, 226)
C_LIGHT_BLUE = RGBColor(224, 242, 254)

FONT_HEADING = 'Arial'
FONT_BODY = 'Arial'
FONT_MONO = 'Consolas'

def create_deck():
    project_root = os.path.dirname(os.path.abspath(__file__))
    template_path = os.path.join(project_root, 'Catenary_Enterprise_Industrial_Twin.pptx')
    prs = Presentation(template_path)
    
    # 1. Delete slide 7 (Instructions slide if present)
    if len(prs.slides) > 6:
        rId = prs.slides._sldIdLst[6].rId
        prs.part.drop_rel(rId)
        del prs.slides._sldIdLst[6]
        print("Deleted instruction Slide 7.")

    # Helper to clean non-template shapes on slide (purging competition artwork)
    def clean_slide_body(slide, keep_types=[13]):
        shapes_to_remove = []
        for s in slide.shapes:
            # Purge competition logos and headers
            if s.shape_type == 13: # Picture
                if "Picture 1" in s.name or "Picture 4" in s.name or "Logo" in s.name:
                    shapes_to_remove.append(s)
                continue
            if s.name.startswith("Rectangle 8") or s.name.startswith("Rectangle 9") or s.name.startswith("Rectangle 24"):
                continue
            if "Placeholder" in s.name and ("Footer" in s.name or "Slide Number" in s.name or "Title" in s.name or "Subtitle" in s.name):
                continue
            if s.name.startswith("Oval"): # Version oval
                continue
            shapes_to_remove.append(s)
            
        for s in shapes_to_remove:
            sp = s._element
            sp.getparent().remove(sp)

    # Helper: add styled card container
    def add_card(slide, left, top, width, height, bg_color=C_CARD_BG, border_color=C_BORDER, line_width=1.0):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = bg_color
        shape.line.color.rgb = border_color
        shape.line.width = Pt(line_width)
        return shape

    # Helper: format footer
    def format_footer(slide, slide_num, team_text="CATENARY | v2.0"):
        for s in slide.shapes:
            if "Footer" in s.name and s.has_text_frame:
                s.text_frame.word_wrap = True
                p = s.text_frame.paragraphs[0]
                p.text = f"Catenary Enterprise Industrial Twin | Oil India Limited (Baghewala Well #14)" if slide_num > 1 else "CATENARY ENTERPRISE INDUSTRIAL TWIN • WELL-TO-SURFACE CYBERNETICS"
                p.font.name = FONT_BODY
                p.font.size = Pt(9.5)
                p.font.color.rgb = C_SLATE
            elif "Slide Number" in s.name and s.has_text_frame:
                p = s.text_frame.paragraphs[0]
                p.text = str(slide_num)
                p.font.name = FONT_BODY
                p.font.size = Pt(10)
                p.font.color.rgb = C_NAVY
            elif s.name.startswith("Oval") and s.has_text_frame:
                s.fill.solid()
                s.fill.fore_color.rgb = C_NAVY
                s.line.color.rgb = C_BLUE
                s.line.width = Pt(1.5)
                p = s.text_frame.paragraphs[0]
                p.text = "CATENARY\nv2.0"
                p.font.name = FONT_HEADING
                p.font.size = Pt(8.5)
                p.font.bold = True
                p.font.color.rgb = C_WHITE
                p.alignment = PP_ALIGN.CENTER

    # Helper: add header & subtitle
    def set_slide_header(slide, title_text, subtitle_text):
        for s in slide.shapes:
            if ("Title 1" in s.name or "Title 7" in s.name) and s.has_text_frame:
                tf = s.text_frame
                tf.clear()
                p = tf.paragraphs[0]
                p.text = title_text
                p.font.name = FONT_HEADING
                p.font.size = Pt(20)
                p.font.bold = True
                p.font.color.rgb = C_NAVY
                
                if subtitle_text:
                    p2 = tf.add_paragraph()
                    p2.text = subtitle_text
                    p2.font.name = FONT_BODY
                    p2.font.size = Pt(11)
                    p2.font.color.rgb = C_BLUE
                    p2.font.bold = True

    # =========================================================================
    # SLIDE 1: TITLE PAGE & TARGET ASSET SNAPSHOT
    # =========================================================================
    print("Formatting Slide 1...")
    s1 = prs.slides[0]
    # Remove old textboxes and competition graphics
    for s in list(s1.shapes):
        if s.name in ["TextBox 9", "Subtitle 3", "Title 7", "Freeform: Shape 26", "Picture 4", "Picture 1"]:
            sp = s._element
            sp.getparent().remove(sp)

    # Top Header Ribbon Banner
    top_banner = add_card(s1, Inches(0.4), Inches(0.35), Inches(9.8), Inches(1.35), bg_color=C_CARD_BG, border_color=C_BLUE, line_width=1.5)
    tb_tf = top_banner.text_frame
    tb_tf.word_wrap = True
    tb_tf.margin_left = Inches(0.2)
    tb_tf.margin_top = Inches(0.12)
    
    p0 = tb_tf.paragraphs[0]
    p0.text = "CATENARY ENTERPRISE INDUSTRIAL TWIN • WELL-TO-SURFACE CYBERNETICS"
    p0.font.name = FONT_HEADING
    p0.font.size = Pt(9.5)
    p0.font.bold = True
    p0.font.color.rgb = C_BLUE

    p1 = tb_tf.add_paragraph()
    p1.text = "WELL-TO-SURFACE PHYSICS-INFORMED DIGITAL TWIN"
    p1.font.name = FONT_HEADING
    p1.font.size = Pt(17)
    p1.font.bold = True
    p1.font.color.rgb = C_NAVY

    p2 = tb_tf.add_paragraph()
    p2.text = "Anticipatory Multi-Timescale Optimization for CSS + SRP Heavy-Oil Operations"
    p2.font.name = FONT_BODY
    p2.font.size = Pt(10.5)
    p2.font.color.rgb = C_SLATE

    # Left Column: Project Context Card
    left_card = add_card(s1, Inches(0.4), Inches(1.85), Inches(6.0), Inches(4.7), bg_color=C_WHITE, border_color=C_BORDER, line_width=1.2)
    lc_tf = left_card.text_frame
    lc_tf.word_wrap = True
    lc_tf.margin_left = Inches(0.25)
    lc_tf.margin_top = Inches(0.2)

    p = lc_tf.paragraphs[0]
    p.text = "SYSTEM METADATA & DEPLOYMENT CONTEXT"
    p.font.name = FONT_HEADING
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = C_NAVY

    meta_items = [
        ("System Identifier", "OIL-BAGHEWALA-EOR-V2", C_BLUE, True),
        ("Asset Specification", "Well-to-Surface Digital Twin for CSS + SRP Optimization\n(Baghewala Heavy Oil Field, Oil India Limited)", C_NAVY, False),
        ("Domain Theme", "Autonomous Industrial Cybernetics / Energy Optimization", C_SLATE, False),
        ("Deployment Architecture", "Production Edge SCADA & Physics-Informed Engine", C_SLATE, False),
        ("Asset Attribution", "Asset: Well #14, Baghewala Heavy Oil Asset, Bikaner-Nagaur Basin, Rajasthan | Operator: Oil India Limited", C_NAVY, True),
        ("System Platform", "Catenary Enterprise Industrial Twin", C_EMERALD, True),
        ("Release Version", "v2.0 (Enterprise Production)", C_EMERALD, True)
    ]

    for label, val, color, bold in meta_items:
        p_lbl = lc_tf.add_paragraph()
        p_lbl.text = f"{label}:"
        p_lbl.font.name = FONT_HEADING
        p_lbl.font.size = Pt(9.5)
        p_lbl.font.bold = True
        p_lbl.font.color.rgb = C_SLATE
        p_lbl.space_before = Pt(4)
        
        p_val = lc_tf.add_paragraph()
        p_val.text = val
        p_val.font.name = FONT_BODY
        p_val.font.size = Pt(10.5)
        p_val.font.bold = bold
        p_val.font.color.rgb = color
        p_val.space_after = Pt(2)

    # Right Column: Target Asset Snapshot Card
    right_card = add_card(s1, Inches(6.6), Inches(1.85), Inches(6.35), Inches(4.7), bg_color=C_CARD_BG, border_color=C_BLUE, line_width=1.5)
    rc_tf = right_card.text_frame
    rc_tf.word_wrap = True
    rc_tf.margin_left = Inches(0.25)
    rc_tf.margin_top = Inches(0.2)

    p = rc_tf.paragraphs[0]
    p.text = "TARGET ASSET SNAPSHOT — BAGHEWALA FIELD"
    p.font.name = FONT_HEADING
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = C_NAVY

    asset_specs = [
        ("Field & Operator", "Baghewala Heavy Oil Field, Rajasthan • Oil India Limited"),
        ("Geological Target", "Jodhpur Sandstone • 1,150 m True Vertical Depth (TVD)"),
        ("Fluid Properties", "14°–19° API • 10,000–13,000 cP @ 48°C native (9 cP @ 260°C)"),
        ("Sector Scale", "35 wells drilled • 23-well active CSS sector (~1,202 BOPD)"),
        ("Artificial Lift", "Cyclic Steam Stimulation (CSS) + Sucker Rod Pump (SRP)"),
        ("Direct Failure Risk", "₹8.5 Lakhs direct rig-repair cost per parted rod string"),
        ("Demonstrated Capability", "4.2 h Warning Lead Time | 12 h Horizon | F_axial >= +0.50 kN")
    ]

    for label, val in asset_specs:
        p_item = rc_tf.add_paragraph()
        p_item.text = f"• {label}: "
        p_item.font.name = FONT_HEADING
        p_item.font.size = Pt(9.5)
        p_item.font.bold = True
        p_item.font.color.rgb = C_SLATE
        p_item.space_before = Pt(3)

        run = p_item.add_run()
        run.text = val
        run.font.name = FONT_BODY
        run.font.size = Pt(9.5)
        run.font.bold = False
        run.font.color.rgb = C_NAVY

    # Capability Badges strip
    badges = ["4.2 h WARNING LEAD", "1D WAVE PDE CORE", "BOBERG-LANTZ THERMAL", "FAST MPC 1.0-6.0 SPM"]
    b_width = Inches(1.4)
    b_start = Inches(6.7)
    for idx, b_text in enumerate(badges):
        b_shape = add_card(s1, b_start + Inches(idx * 1.48), Inches(6.05), Inches(1.42), Inches(0.35), bg_color=C_NAVY, border_color=C_BLUE, line_width=1.0)
        btf = b_shape.text_frame
        btf.word_wrap = True
        btf.margin_top = Inches(0.04)
        bp = btf.paragraphs[0]
        bp.text = b_text
        bp.font.name = FONT_HEADING
        bp.font.size = Pt(7.2)
        bp.font.bold = True
        bp.font.color.rgb = C_WHITE
        bp.alignment = PP_ALIGN.CENTER

    format_footer(s1, 1)

    # =========================================================================
    # SLIDE 2: PROPOSED SOLUTION & PROTOTYPE SHOWCASE
    # =========================================================================
    print("Formatting Slide 2...")
    s2 = prs.slides[1]
    clean_slide_body(s2)
    set_slide_header(s2, "PROPOSED SOLUTION: COUPLED DIGITAL TWIN", "Eliminating the Heavy-Oil Viscosity Trap via Anticipatory Thermal-Mechanical Coupling")
    format_footer(s2, 2)

    # Top Problem Callout Strip
    prob_box = add_card(s2, Inches(0.4), Inches(1.15), Inches(12.5), Inches(0.65), bg_color=C_LIGHT_RED, border_color=C_CRIMSON, line_width=1.2)
    pbtf = prob_box.text_frame
    pbtf.word_wrap = True
    pbtf.margin_left = Inches(0.15)
    pbtf.margin_top = Inches(0.05)
    p = pbtf.paragraphs[0]
    p.text = "THE VISCOSITY TRAP: "
    p.font.name = FONT_HEADING
    p.font.size = Pt(9.5)
    p.font.bold = True
    p.font.color.rgb = C_CRIMSON
    r = p.add_run()
    r.text = "Post-steam cooling triggers exponential Arrhenius viscosity surge (9 cP → 12,000 cP), causing severe downstroke drag, compressive rod float (F_axial < 0 kN), buckling, and violent parting — exposing the 23-well sector to ₹4.01 Cr annual direct rig repairs."
    r.font.name = FONT_BODY
    r.font.size = Pt(8.8)
    r.font.color.rgb = C_NAVY

    # Left Box: Live UI Prototype Console
    ui_img_path = 'presentation_assets/slide2_console_ui.png'
    s2.shapes.add_picture(ui_img_path, Inches(0.4), Inches(1.88), Inches(6.2), Inches(3.95))

    # Right Box: Four Innovation Pillars
    right_pillar_box = add_card(s2, Inches(6.75), Inches(1.88), Inches(6.15), Inches(3.95), bg_color=C_WHITE, border_color=C_BORDER, line_width=1.2)
    rtf = right_pillar_box.text_frame
    rtf.word_wrap = True
    rtf.margin_left = Inches(0.2)
    rtf.margin_top = Inches(0.12)

    p = rtf.paragraphs[0]
    p.text = "FOUR CORE INNOVATION PILLARS"
    p.font.name = FONT_HEADING
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = C_NAVY

    pillars = [
        ("1. Anticipatory Thermal Feedforward", "Boberg-Lantz analytical decay + quad integration maps cooling into live VFD speed dispatches with 4.2 h early warning lead time."),
        ("2. Hard Structural Tension Envelopes", "Enforces downhole rod tension >= +0.50 kN floor & PPRL <= 90% via KKT-constrained fast-loop MPC optimizer (1.0 - 6.0 SPM)."),
        ("3. Multi-Timescale Cybernetic Control", "Unifies slow multi-week steam cooldown dynamics with sub-second (1.56 ms CFL) 1D elastodynamic wave mechanics."),
        ("4. Cryptographic Provenance & SCADA", "SHA-256 tamper-evident logs tagging [measured], [model], and [synthetic] data with Modbus RTU/TCP & MQTT integration.")
    ]

    for title, body in pillars:
        p_t = rtf.add_paragraph()
        p_t.text = title
        p_t.font.name = FONT_HEADING
        p_t.font.size = Pt(9.2)
        p_t.font.bold = True
        p_t.font.color.rgb = C_BLUE
        p_t.space_before = Pt(3)

        p_b = rtf.add_paragraph()
        p_b.text = body
        p_b.font.name = FONT_BODY
        p_b.font.size = Pt(8.5)
        p_b.font.color.rgb = C_SLATE
        p_b.space_after = Pt(2)

    # Bottom Result Ribbon
    res_box = add_card(s2, Inches(0.4), Inches(5.92), Inches(12.5), Inches(0.55), bg_color=C_LIGHT_GREEN, border_color=C_EMERALD, line_width=1.5)
    rbtf = res_box.text_frame
    rbtf.word_wrap = True
    rbtf.margin_left = Inches(0.15)
    rbtf.margin_top = Inches(0.08)
    p = rbtf.paragraphs[0]
    p.text = "★ DEMONSTRATED RESULT (4.2h INTERCEPT): "
    p.font.name = FONT_HEADING
    p.font.size = Pt(10)
    p.font.bold = True
    p.font.color.rgb = C_EMERALD
    r = p.add_run()
    r.text = "Fast MPC throttles pump speed (4.7 → 2.8 SPM) • Continuous positive rod tension held at +0.65 kN (Floor >= +0.50 kN) • ZERO rod float, ZERO buckling!"
    r.font.name = FONT_BODY
    r.font.size = Pt(9.2)
    r.font.bold = True
    r.font.color.rgb = C_NAVY

    # =========================================================================
    # SLIDE 3: TECHNICAL APPROACH & SYSTEM ARCHITECTURE
    # =========================================================================
    print("Formatting Slide 3...")
    s3 = prs.slides[2]
    clean_slide_body(s3)
    set_slide_header(s3, "TECHNICAL APPROACH & SYSTEM ARCHITECTURE", "5-Module Physics-to-Control Core & Clean-Room Industrial Edge Stack")
    format_footer(s3, 3)

    # Left Column: 5-Module Physics Engine Architecture diagram
    arch_img_path = 'presentation_assets/slide3_architecture.png'
    s3.shapes.add_picture(arch_img_path, Inches(0.4), Inches(1.15), Inches(7.0), Inches(5.35))

    # Right Column: Tech Stack & Compliance Cards
    # Card 1: Clean-Room Tech Stack
    stack_card = add_card(s3, Inches(7.55), Inches(1.15), Inches(5.35), Inches(3.0), bg_color=C_WHITE, border_color=C_BORDER, line_width=1.2)
    stf = stack_card.text_frame
    stf.word_wrap = True
    stf.margin_left = Inches(0.2)
    stf.margin_top = Inches(0.12)

    p = stf.paragraphs[0]
    p.text = "CLEAN-ROOM OPEN-SOURCE TECH STACK"
    p.font.name = FONT_HEADING
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = C_NAVY

    stack_items = [
        ("Physics Solvers", "Pure Clean-Room Python, NumPy, SciPy (Bessel quad)"),
        ("Dynamic Optimization", "CasADi / OSQP (Fast-Loop MPC with hard KKT constraints)"),
        ("Wave Elastodynamics", "Gibbs 1D damped PDE solver with CFL dt <= 1.56 ms"),
        ("Data Contract", "Pydantic schema validation + SHA-256 cryptographic chain"),
        ("Edge SCADA & Telemetry", "pymodbus (Modbus RTU/TCP) + Eclipse Mosquitto (MQTT)")
    ]

    for cat, desc in stack_items:
        p_item = stf.add_paragraph()
        p_item.text = f"• {cat}: "
        p_item.font.name = FONT_HEADING
        p_item.font.size = Pt(8.8)
        p_item.font.bold = True
        p_item.font.color.rgb = C_BLUE
        p_item.space_before = Pt(3)

        run = p_item.add_run()
        run.text = desc
        run.font.name = FONT_BODY
        run.font.size = Pt(8.5)
        run.font.color.rgb = C_SLATE

    # Card 2: Enterprise Specifications & Compliance
    eval_card = add_card(s3, Inches(7.55), Inches(4.25), Inches(5.35), Inches(2.25), bg_color=C_LIGHT_GREEN, border_color=C_EMERALD, line_width=1.5)
    etf = eval_card.text_frame
    etf.word_wrap = True
    etf.margin_left = Inches(0.2)
    etf.margin_top = Inches(0.12)

    p = etf.paragraphs[0]
    p.text = "ENTERPRISE INDUSTRIAL SPECIFICATIONS & COMPLIANCE"
    p.font.name = FONT_HEADING
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = C_EMERALD

    compliances = [
        "✓ 100% Free & Open-Source: BSD / MIT / Apache-2.0 licenses only.",
        "✓ ₹0 Cloud API Fees: Runs 100% locally on standard wellsite edge IPCs without active Wi-Fi / cloud lock-in.",
        "✓ Pre-Validated Core: 255/255 unit tests passing with strict CFL acoustic stability (dt <= 1.56 ms)."
    ]

    for comp in compliances:
        p_c = etf.add_paragraph()
        p_c.text = comp
        p_c.font.name = FONT_BODY
        p_c.font.size = Pt(8.8)
        p_c.font.bold = True
        p_c.font.color.rgb = C_NAVY
        p_c.space_before = Pt(3)

    # =========================================================================
    # SLIDE 4: FEASIBILITY, SAFETY MATRIX & RISK MITIGATION
    # =========================================================================
    print("Formatting Slide 4...")
    s4 = prs.slides[3]
    clean_slide_body(s4)
    set_slide_header(s4, "FEASIBILITY, SAFETY MATRIX & RISK MITIGATION", "Deterministic 4-Level Supervisory Safety State Machine & Operational Resilience")
    format_footer(s4, 4)

    # Top KPI Badges (3 horizontal cards)
    top_kpis = [
        ("✓ 255 / 255 TESTS PASSING", "100% Core Test Coverage", C_EMERALD, C_LIGHT_GREEN),
        ("✓ UNIVERSAL CSV ADAPTER", "Auto-Mapping Field Heuristics", C_BLUE, C_LIGHT_BLUE),
        ("✓ WELLSITE EDGE IPC READY", "Offline Modbus/MQTT Execution", C_NAVY, C_CARD_BG)
    ]
    for i, (t_kpi, s_kpi, fg, bg) in enumerate(top_kpis):
        k_card = add_card(s4, Inches(0.4 + i * 4.2), Inches(1.15), Inches(4.0), Inches(0.65), bg_color=bg, border_color=fg, line_width=1.2)
        ktf = k_card.text_frame
        ktf.word_wrap = True
        ktf.margin_top = Inches(0.06)
        p1 = ktf.paragraphs[0]
        p1.text = t_kpi
        p1.font.name = FONT_HEADING
        p1.font.size = Pt(9.5)
        p1.font.bold = True
        p1.font.color.rgb = fg
        p1.alignment = PP_ALIGN.CENTER
        
        p2 = ktf.add_paragraph()
        p2.text = s_kpi
        p2.font.name = FONT_BODY
        p2.font.size = Pt(8.2)
        p2.font.color.rgb = C_SLATE
        p2.alignment = PP_ALIGN.CENTER

    # State Machine Flowchart graphic
    sm_img_path = 'presentation_assets/slide4_statemachine.png'
    s4.shapes.add_picture(sm_img_path, Inches(0.4), Inches(1.88), Inches(12.5), Inches(2.35))

    # Bottom Section: Risk Mitigation Matrix (3 Columns)
    risks = [
        ("RISK: Boberg-Lantz 42% Late Bias",
         "Analytical Boberg-Lantz overpredicts late-time reservoir temperature by up to 42% at ~300 days.",
         "MITIGATION: Dynamic k(t) Scaling",
         "Calibrates dynamic scaling factor k(t) smoothly from 1.0 down to 0.7042, clamping to native T_R.",
         C_BLUE, C_LIGHT_BLUE),
         
        ("RISK: Non-Linear Post-Buckling",
         "Compressive loads trigger non-linear sinusoidal rod contact against tubing walls, causing wear & parting.",
         "MITIGATION: Tensile Barrier Constraint",
         "Enforces hard lower bound F_axial >= +0.50 kN in MPC, completely avoiding compressive buckling regime.",
         C_EMERALD, C_LIGHT_GREEN),
         
        ("RISK: Desert Telemetry Gaps",
         "Harsh Thar Desert conditions cause intermittent Modbus packet loss and sensor latency spikes.",
         "MITIGATION: Auto Level-2 Fallback",
         "Bounded interpolation (<=3 samples) + automated deterministic 3-stroke ramp down to safe 2.0 SPM fallback.",
         C_AMBER, C_CARD_BG)
    ]

    for i, (r_title, r_desc, m_title, m_desc, border_c, bg_c) in enumerate(risks):
        r_card = add_card(s4, Inches(0.4 + i * 4.2), Inches(4.30), Inches(4.0), Inches(2.15), bg_color=C_WHITE, border_color=border_c, line_width=1.3)
        rtf = r_card.text_frame
        rtf.word_wrap = True
        rtf.margin_left = Inches(0.15)
        rtf.margin_top = Inches(0.1)

        p = rtf.paragraphs[0]
        p.text = r_title
        p.font.name = FONT_HEADING
        p.font.size = Pt(8.8)
        p.font.bold = True
        p.font.color.rgb = C_CRIMSON

        p2 = rtf.add_paragraph()
        p2.text = r_desc
        p2.font.name = FONT_BODY
        p2.font.size = Pt(8.0)
        p2.font.color.rgb = C_SLATE
        p2.space_after = Pt(2)

        p3 = rtf.add_paragraph()
        p3.text = m_title
        p3.font.name = FONT_HEADING
        p3.font.size = Pt(8.8)
        p3.font.bold = True
        p3.font.color.rgb = border_c

        p4 = rtf.add_paragraph()
        p4.text = m_desc
        p4.font.name = FONT_BODY
        p4.font.size = Pt(8.0)
        p4.font.color.rgb = C_NAVY

    # =========================================================================
    # SLIDE 5: OPERATIONAL IMPACT & TWO-TIER ASSET ECONOMICS
    # =========================================================================
    print("Formatting Slide 5...")
    s5 = prs.slides[4]
    clean_slide_body(s5)
    set_slide_header(s5, "OPERATIONAL IMPACT & QUANTIFIED ASSET ROI", "Field-Wide Economic Recovery Model for Baghewala's 23 Active Wells")
    format_footer(s5, 5)

    # Left Column: Operational Comparison Table
    # Table container
    left_tbl_card = add_card(s5, Inches(0.4), Inches(1.15), Inches(6.0), Inches(4.6), bg_color=C_WHITE, border_color=C_BORDER, line_width=1.2)
    ltf = left_tbl_card.text_frame
    ltf.word_wrap = True
    ltf.margin_left = Inches(0.2)
    ltf.margin_top = Inches(0.12)
    p = ltf.paragraphs[0]
    p.text = "OPERATIONAL PERFORMANCE COMPARISON"
    p.font.name = FONT_HEADING
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = C_NAVY

    # Create table inside
    table_shape = s5.shapes.add_table(6, 4, Inches(0.55), Inches(1.60), Inches(5.7), Inches(3.95))
    tbl = table_shape.table
    tbl.columns[0].width = Inches(1.6)
    tbl.columns[1].width = Inches(1.3)
    tbl.columns[2].width = Inches(1.3)
    tbl.columns[3].width = Inches(1.5)

    headers = ["Operational Metric", "Reactive Baseline", "Digital Twin Target", "Field Benefit"]
    for col_idx, h in enumerate(headers):
        cell = tbl.cell(0, col_idx)
        cell.fill.solid()
        cell.fill.fore_color.rgb = C_NAVY
        p = cell.text_frame.paragraphs[0]
        p.text = h
        p.font.name = FONT_HEADING
        p.font.size = Pt(8.0)
        p.font.bold = True
        p.font.color.rgb = C_WHITE
        p.alignment = PP_ALIGN.CENTER

    table_data = [
        ("Warning Lead Time", "0 Hours (Reactive)", "4.2 Hours Early", "Preemptive Throttling"),
        ("Degraded Drag Strokes", "~1,280 Strokes", "<= 100 Strokes", "92% Reduction in Drag"),
        ("Downhole Impact Strokes", "~640 Strokes", "0 - 50 Strokes", "Eliminates Collisions"),
        ("Min Rod Tension (F_min)", "-1.80 kN (Float)", ">= +0.50 kN (Safe)", "Eliminates Buckling"),
        ("Peak Load (PPRL)", "115% Rating (Overload)", "<= 90% Structural", "Preserves Fatigue Life")
    ]

    for row_idx, row in enumerate(table_data):
        for col_idx, val in enumerate(row):
            cell = tbl.cell(row_idx + 1, col_idx)
            cell.fill.solid()
            cell.fill.fore_color.rgb = C_CARD_BG if row_idx % 2 == 0 else C_WHITE
            p = cell.text_frame.paragraphs[0]
            p.text = val
            p.font.name = FONT_BODY
            p.font.size = Pt(7.6)
            if col_idx == 0:
                p.font.bold = True
                p.font.color.rgb = C_NAVY
            elif col_idx == 1:
                p.font.color.rgb = C_CRIMSON
            elif col_idx == 2:
                p.font.bold = True
                p.font.color.rgb = C_EMERALD
            else:
                p.font.color.rgb = C_BLUE
            p.alignment = PP_ALIGN.CENTER if col_idx > 0 else PP_ALIGN.LEFT

    # Right Column: Two-Tier Valuation & Financial Waterfall
    right_val_card = add_card(s5, Inches(6.6), Inches(1.15), Inches(6.35), Inches(4.6), bg_color=C_WHITE, border_color=C_BORDER, line_width=1.2)
    rtf = right_val_card.text_frame
    rtf.word_wrap = True
    rtf.margin_left = Inches(0.2)
    rtf.margin_top = Inches(0.12)
    p = rtf.paragraphs[0]
    p.text = "BAGHEWALA 23-WELL SECTOR ANNUAL ASSET VALUATION"
    p.font.name = FONT_HEADING
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = C_NAVY

    # Embedded Waterfall Chart
    wf_img_path = 'presentation_assets/slide5_waterfall.png'
    s5.shapes.add_picture(wf_img_path, Inches(6.75), Inches(1.55), Inches(6.05), Inches(2.75))

    # Valuation Breakdown Card
    val_strip = add_card(s5, Inches(6.75), Inches(4.35), Inches(6.05), Inches(1.3), bg_color=C_CARD_BG, border_color=C_BLUE, line_width=1.2)
    vtf = val_strip.text_frame
    vtf.word_wrap = True
    vtf.margin_left = Inches(0.15)
    vtf.margin_top = Inches(0.08)

    p1 = vtf.paragraphs[0]
    p1.text = "• TIER-1 GUARANTEED OPEX FLOOR: ₹4.31 CRORE / YEAR"
    p1.font.name = FONT_HEADING
    p1.font.size = Pt(8.8)
    p1.font.bold = True
    p1.font.color.rgb = C_EMERALD
    r1 = p1.add_run()
    r1.text = " (47.15 failures avoided × ₹8.5L = ₹4.01 Cr + 402,960 kWh saved = ₹30.2L)"
    r1.font.name = FONT_BODY
    r1.font.size = Pt(7.8)
    r1.font.color.rgb = C_SLATE

    p2 = vtf.add_paragraph()
    p2.text = "• TIER-2 TOTAL FIELD POTENTIAL: ₹16.4 – ₹29.4 CRORE / YEAR"
    p2.font.name = FONT_HEADING
    p2.font.size = Pt(8.8)
    p2.font.bold = True
    p2.font.color.rgb = C_NAVY
    r2 = p2.add_run()
    r2.text = " (35,259 bbl deferred oil recovered @ $75/bbl, ₹95 FX)"
    r2.font.name = FONT_BODY
    r2.font.size = Pt(7.8)
    r2.font.color.rgb = C_SLATE

    # Bottom Beyond Economics Strip
    beyond_box = add_card(s5, Inches(0.4), Inches(5.85), Inches(12.55), Inches(0.6), bg_color=C_CARD_BG, border_color=C_EMERALD, line_width=1.2)
    btf = beyond_box.text_frame
    btf.word_wrap = True
    btf.margin_left = Inches(0.15)
    btf.margin_top = Inches(0.08)
    p = btf.paragraphs[0]
    p.text = "🌱 ESG & NATIONAL IMPACT: "
    p.font.name = FONT_HEADING
    p.font.size = Pt(9.2)
    p.font.bold = True
    p.font.color.rgb = C_EMERALD
    r = p.add_run()
    r.text = "293 t CO2/yr Avoided (0.727 kg/kWh CEA grid factor) • Pioneer Heavy-Oil Digital Twin for Atmanirbhar Bharat • Direct Scale to ONGC Gujarat Assets"
    r.font.name = FONT_BODY
    r.font.size = Pt(8.5)
    r.font.color.rgb = C_NAVY

    # =========================================================================
    # SLIDE 6: RESEARCH, CITATIONS & THE PILOT ASK
    # =========================================================================
    print("Formatting Slide 6...")
    s6 = prs.slides[5]
    clean_slide_body(s6)
    set_slide_header(s6, "ACADEMIC FOUNDATIONS & DEPLOYMENT PILOT", "Peer-Reviewed Scientific Foundations & Low-Risk 60-Day Field Validation")
    format_footer(s6, 6)

    # Left Column: Primary Literature Card
    lit_card = add_card(s6, Inches(0.4), Inches(1.15), Inches(6.0), Inches(5.30), bg_color=C_WHITE, border_color=C_BORDER, line_width=1.2)
    ltf = lit_card.text_frame
    ltf.word_wrap = True
    ltf.margin_left = Inches(0.2)
    ltf.margin_top = Inches(0.12)

    p = ltf.paragraphs[0]
    p.text = "PEER-REVIEWED SCIENTIFIC LITERATURE"
    p.font.name = FONT_HEADING
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = C_NAVY

    citations = [
        ("Boberg, T. C., & Lantz, R. B. (1966, SPE-1578-PA)", "Calculation of the Production Rate of a Project by Steam Stimulation (Analytical thermal decay foundation)."),
        ("Safari, M., et al. (2020, Petroleum)", "Analytical equations for cyclic steam stimulation & 42% late-time temperature bias validation benchmark."),
        ("Gibbs, S. G. (1963, SPE-588-PA)", "Predicting the Behavior of Sucker-Rod Pumping Systems (1D hyperbolic damped wave PDE mechanics)."),
        ("Hansen, J., et al. (2019, BYU-PRISM)", "Real-time artificial lift wave mechanics & physics-based diagnostic modeling benchmark."),
        ("Cheng, K., et al. (2020, Sensors)", "8-class dynamometer card condition recognition & downhole load fault identification.")
    ]

    for author, desc in citations:
        p_a = ltf.add_paragraph()
        p_a.text = f"• {author}"
        p_a.font.name = FONT_HEADING
        p_a.font.size = Pt(8.5)
        p_a.font.bold = True
        p_a.font.color.rgb = C_BLUE
        p_a.space_before = Pt(3)

        p_d = ltf.add_paragraph()
        p_d.text = f"  {desc}"
        p_d.font.name = FONT_BODY
        p_d.font.size = Pt(7.8)
        p_d.font.color.rgb = C_SLATE

    # Data Provenance Box at bottom of left column
    prov_box = add_card(s6, Inches(0.6), Inches(4.75), Inches(5.6), Inches(1.5), bg_color=C_CARD_BG, border_color=C_BLUE, line_width=1.0)
    ptf = prov_box.text_frame
    ptf.word_wrap = True
    ptf.margin_left = Inches(0.15)
    ptf.margin_top = Inches(0.08)

    p = ptf.paragraphs[0]
    p.text = "DATA REALITY & FIELD INTEGRATION:"
    p.font.name = FONT_HEADING
    p.font.size = Pt(8.8)
    p.font.bold = True
    p.font.color.rgb = C_NAVY

    p2 = ptf.add_paragraph()
    p2.text = "Universal regex auto-mapping CSV adapter ingests raw field telemetry logs in minutes without requiring live network reconfiguration. Tested against 23 well profiles."
    p2.font.name = FONT_BODY
    p2.font.size = Pt(7.8)
    p2.font.color.rgb = C_SLATE

    # Right Column: TRL Ladder & The Pilot Ask
    # TRL Graphic
    trl_img_path = 'presentation_assets/slide6_trl.png'
    s6.shapes.add_picture(trl_img_path, Inches(6.6), Inches(1.15), Inches(6.35), Inches(2.95))

    # The Pilot Ask Callout Card
    ask_card = add_card(s6, Inches(6.6), Inches(4.20), Inches(6.35), Inches(2.25), bg_color=C_LIGHT_GREEN, border_color=C_EMERALD, line_width=1.8)
    atf = ask_card.text_frame
    atf.word_wrap = True
    atf.margin_left = Inches(0.2)
    atf.margin_top = Inches(0.12)

    p = atf.paragraphs[0]
    p.text = "THE ASK: 60-DAY SHADOW-MODE PILOT (OIL INDIA LIMITED)"
    p.font.name = FONT_HEADING
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = C_EMERALD

    ask_points = [
        ("Non-Actuating Shadow Deployment", "Installs read-only alongside existing RTUs with ZERO operational or hardware risk."),
        ("Objective & Calibration", "Validates and fine-tunes dynamic k(t) & Couette drag parameters against live Baghewala logs."),
        ("Technology Readiness Advancement", "Advances proven digital twin core from TRL 4-5 to TRL 7 commercial fleet rollout across 23 wells.")
    ]

    for title, desc in ask_points:
        p_item = atf.add_paragraph()
        p_item.text = f"★ {title}: "
        p_item.font.name = FONT_HEADING
        p_item.font.size = Pt(8.8)
        p_item.font.bold = True
        p_item.font.color.rgb = C_NAVY
        p_item.space_before = Pt(2)

        run = p_item.add_run()
        run.text = desc
        run.font.name = FONT_BODY
        run.font.size = Pt(8.0)
        run.font.color.rgb = C_SLATE

    # Save final presentation
    output_path = os.path.join(project_root, 'Catenary_Enterprise_Industrial_Twin.pptx')
    prs.save(output_path)
    print(f"Presentation saved to: {output_path}")

if __name__ == '__main__':
    create_deck()
